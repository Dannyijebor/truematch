import io
import os
import json
import httpx
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")


def extract_text(file_bytes: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    if name.endswith(".docx"):
        return _extract_docx(file_bytes)
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_pdf(file_bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


STRUCTURE_SYSTEM = """You are a precise resume parser. Extract ONLY what is present.
NEVER invent employers, dates, degrees, certifications, or metrics.
If a field is missing, use null or empty.

Return ONLY valid JSON:
{
  "name": "",
  "email": null,
  "phone": null,
  "location": null,
  "headline": "",
  "summary": "",
  "experience": [{"company":"","title":"","start":null,"end":null,"location":null,"bullets":[]}],
  "education": [{"school":"","degree":"","field":"","start":"","end":""}],
  "skills": [],
  "projects": [{"name":"","description":"","tech":[]}],
  "certifications": [],
  "languages": []
}
"""


def structure_resume(raw_text: str) -> dict:
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise RuntimeError("GROQ_API_KEY missing")
    body = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": STRUCTURE_SYSTEM},
            {"role": "user", "content": "RESUME TEXT:\n\n" + raw_text[:15000]},
        ],
        "temperature": 0.1,
        "max_tokens": 4000,
        "response_format": {"type": "json_object"},
    }
    r = httpx.post(GROQ_URL, headers={
        "Authorization": f"Bearer {key}", "Content-Type": "application/json"
    }, json=body, timeout=90)
    r.raise_for_status()
    return json.loads(r.json()["choices"][0]["message"]["content"])


TAILOR_SYSTEM = """You are a resume coach. Given a CANDIDATE's structured resume and a JOB,
produce an improved version better aligned to the job.

ABSOLUTE RULES:
1. NEVER invent employers, dates, degrees, certifications, or skills.
2. Every fact must exist in the CANDIDATE input.
3. Rewrite summary and bullets to use language from the job where truthful.
4. Quantify ONLY if the source already has numbers.
5. Bullets must be 18-28 words, start with a strong verb.
6. Return ONLY valid JSON in the SAME schema as the CANDIDATE.
"""


def tailor_resume(structured: dict, job, profile=None) -> dict:
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise RuntimeError("GROQ_API_KEY missing")

    job_blob = f"""Title: {job.title}
Company: {job.company.name if job.company else 'N/A'}
Location: {job.location or 'N/A'}

Description:
{(job.description_text or '')[:6000]}
"""
    user_prompt = (
        "CANDIDATE:\n" + json.dumps(structured, ensure_ascii=False)[:12000]
        + "\n\nJOB:\n" + job_blob
        + "\n\nReturn the tailored structured resume JSON."
    )
    body = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": TAILOR_SYSTEM},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.3,
        "max_tokens": 4000,
        "response_format": {"type": "json_object"},
    }
    r = httpx.post(GROQ_URL, headers={
        "Authorization": f"Bearer {key}", "Content-Type": "application/json"
    }, json=body, timeout=90)
    r.raise_for_status()
    return json.loads(r.json()["choices"][0]["message"]["content"])


QUESTIONS_SYSTEM = """You create 4 short, targeted questions to help tailor a resume to a specific job.
The questions must help the candidate reveal truthful, job-relevant experience.

Return ONLY valid JSON:
{
  "questions": [
    {"id": "kubernetes_years", "question": "How many years have you worked with Kubernetes?", "type": "short"},
    {"id": "leadership_example", "question": "Describe one team you led and the measurable outcome.", "type": "long"}
  ]
}
"""


def generate_job_questions(job) -> list[dict]:
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise RuntimeError("GROQ_API_KEY missing")

    blob = f"{job.title}\n{(job.description_text or '')[:5000]}"
    body = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": QUESTIONS_SYSTEM},
            {"role": "user", "content": "JOB:\n" + blob},
        ],
        "temperature": 0.4,
        "max_tokens": 800,
        "response_format": {"type": "json_object"},
    }
    r = httpx.post(GROQ_URL, headers={
        "Authorization": f"Bearer {key}", "Content-Type": "application/json"
    }, json=body, timeout=60)
    r.raise_for_status()
    data = json.loads(r.json()["choices"][0]["message"]["content"])
    return data.get("questions", [])


def generate_docx(structured: dict) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    def heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)

    name = structured.get("name") or "Your Name"
    p = doc.add_paragraph()
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    bits = [structured.get("email"), structured.get("phone"), structured.get("location")]
    contact = " · ".join([b for b in bits if b])
    if contact:
        p = doc.add_paragraph()
        r = p.add_run(contact)
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    if structured.get("headline"):
        p = doc.add_paragraph()
        r = p.add_run(structured["headline"])
        r.italic = True
        r.font.size = Pt(11)

    if structured.get("summary"):
        heading("Summary")
        doc.add_paragraph(structured["summary"])

    for e in structured.get("experience") or []:
        if not e.get("title") and not e.get("company"):
            continue
        heading("Experience") if (e is (structured.get("experience") or [None])[0]) else None
        p = doc.add_paragraph()
        title = e.get("title") or ""
        comp = e.get("company") or ""
        line = f"{title} — {comp}" if title and comp else (title or comp)
        r = p.add_run(line)
        r.bold = True
        r.font.size = Pt(11)
        meta = " · ".join([x for x in [e.get("start"), e.get("end"), e.get("location")] if x])
        if meta:
            p = doc.add_paragraph()
            r = p.add_run(meta)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        for b in (e.get("bullets") or []):
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(2)

    edu = structured.get("education") or []
    if edu:
        heading("Education")
        for e in edu:
            line = " — ".join([x for x in [e.get("degree"), e.get("field"), e.get("school")] if x])
            if not line:
                continue
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            meta = " · ".join([x for x in [e.get("start"), e.get("end")] if x])
            if meta:
                p = doc.add_paragraph()
                r = p.add_run(meta)
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    skills = structured.get("skills") or []
    if skills:
        heading("Skills")
        doc.add_paragraph(", ".join(skills))

    projs = structured.get("projects") or []
    if projs:
        heading("Projects")
        for pr in projs:
            if not pr.get("name") and not pr.get("description"):
                continue
            p = doc.add_paragraph()
            r = p.add_run(pr.get("name") or "")
            r.bold = True
            if pr.get("description"):
                doc.add_paragraph(pr["description"])
            if pr.get("tech"):
                p = doc.add_paragraph()
                r = p.add_run("Tech: " + ", ".join(pr["tech"]))
                r.font.size = Pt(9)
                r.italic = True

    certs = structured.get("certifications") or []
    if certs:
        heading("Certifications")
        for c in certs:
            doc.add_paragraph(c, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
